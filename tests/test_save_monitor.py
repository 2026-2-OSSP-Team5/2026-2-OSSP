import io
import logging
import queue
import time
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from slideguard.adapters.watching.monitor import SavedFileMonitor


def pptx_bytes(text: str) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2)).text = text
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_baseline_and_identical_bytes_are_not_new_saves(tmp_path, caplog):
    caplog.set_level(logging.INFO, logger="slideguard.watch")
    path = tmp_path / "발표.PPTX"
    original = pptx_bytes("initial")
    path.write_bytes(original)
    saved = queue.Queue()
    monitor = SavedFileMonitor(tmp_path, saved.put, settle_seconds=0.15)
    monitor.start()
    try:
        path.write_bytes(original)
        time.sleep(0.6)
        assert saved.empty()
        path.write_bytes(pptx_bytes("changed"))
        assert saved.get(timeout=5) == path
        time.sleep(0.5)
        assert saved.empty()
        assert sum("SAVE_DETECTED" in r.message for r in caplog.records) == 1
    finally:
        monitor.stop()
    assert not monitor.is_running


def test_partial_zip_and_repeated_writes_emit_only_completed_save(tmp_path):
    path = tmp_path / "partial.pptx"
    saved = queue.Queue()
    monitor = SavedFileMonitor(tmp_path, saved.put, settle_seconds=0.15)
    monitor.start()
    try:
        path.write_bytes(b"incomplete ZIP")
        time.sleep(0.5)
        assert saved.empty()
        completed = pptx_bytes("complete")
        with path.open("wb") as stream:
            stream.write(completed[: len(completed) // 2])
            stream.flush()
            time.sleep(0.1)
            assert saved.empty()
            stream.write(completed[len(completed) // 2 :])
        assert saved.get(timeout=5) == path
        time.sleep(0.5)
        assert saved.empty()
    finally:
        monitor.stop()


def test_atomic_replace_save_is_detected(tmp_path):
    path = tmp_path / "atomic.pptx"
    path.write_bytes(pptx_bytes("old"))
    saved = queue.Queue()
    monitor = SavedFileMonitor(tmp_path, saved.put, settle_seconds=0.15)
    monitor.start()
    try:
        temporary = tmp_path / "ppt-save.tmp"
        temporary.write_bytes(pptx_bytes("new"))
        temporary.replace(path)
        assert saved.get(timeout=5) == path
        time.sleep(0.5)
        assert saved.empty()
    finally:
        monitor.stop()


def test_lock_files_and_other_extensions_are_ignored(tmp_path):
    saved = queue.Queue()
    monitor = SavedFileMonitor(tmp_path, saved.put, settle_seconds=0.15)
    monitor.start()
    try:
        data = pptx_bytes("ignore")
        for name in ("~$deck.pptx", "~scratch.pptx", ".temporary.pptx", "deck.tmp", "deck.ppt"):
            (tmp_path / name).write_bytes(data)
        time.sleep(0.6)
        assert saved.empty()
        valid = tmp_path / "deck.pptx"
        valid.write_bytes(data)
        assert saved.get(timeout=5) == valid
    finally:
        monitor.stop()


@pytest.mark.parametrize("recursive", [False, True])
def test_subfolders_follow_recursive_option(tmp_path, recursive):
    nested = tmp_path / "nested"
    nested.mkdir()
    saved = queue.Queue()
    monitor = SavedFileMonitor(tmp_path, saved.put, recursive=recursive, settle_seconds=0.15)
    monitor.start()
    try:
        path = nested / "lecture.pptx"
        path.write_bytes(pptx_bytes("nested"))
        if recursive:
            assert saved.get(timeout=5) == path
        else:
            time.sleep(0.6)
            assert saved.empty()
    finally:
        monitor.stop()


def test_moved_folder_is_scanned_when_recursive(tmp_path):
    watched = tmp_path / "watched"
    incoming = tmp_path / "incoming"
    watched.mkdir()
    incoming.mkdir()
    (incoming / "moved.pptx").write_bytes(pptx_bytes("move folder"))
    saved = queue.Queue()
    monitor = SavedFileMonitor(watched, saved.put, recursive=True, settle_seconds=0.15)
    monitor.start()
    try:
        destination = watched / "incoming"
        incoming.rename(destination)
        assert saved.get(timeout=5) == destination / "moved.pptx"
    finally:
        monitor.stop()


def test_invalid_file_times_out_then_next_write_can_be_detected(tmp_path, caplog):
    caplog.set_level(logging.WARNING, logger="slideguard.watch")
    saved = queue.Queue()
    monitor = SavedFileMonitor(tmp_path, saved.put, settle_seconds=0.1, retry_seconds=0.3)
    monitor.start()
    try:
        path = tmp_path / "broken.pptx"
        path.write_bytes(b"not a presentation")
        deadline = time.monotonic() + 5
        while not any("SAVE_UNREADABLE" in r.message for r in caplog.records):
            assert time.monotonic() < deadline
            time.sleep(0.05)
        assert saved.empty()
        path.write_bytes(pptx_bytes("fixed"))
        assert saved.get(timeout=5) == path
    finally:
        monitor.stop()


def test_invalid_directory_is_rejected(tmp_path):
    with pytest.raises(NotADirectoryError):
        SavedFileMonitor(tmp_path / "missing").start()


def test_callback_failure_does_not_stop_following_saves(tmp_path, caplog):
    saved = queue.Queue()

    def consume(path: Path):
        saved.put(path)
        raise RuntimeError("consumer failed")

    monitor = SavedFileMonitor(tmp_path, consume, settle_seconds=0.15)
    monitor.start()
    try:
        path = tmp_path / "callback.pptx"
        path.write_bytes(pptx_bytes("one"))
        assert saved.get(timeout=5) == path
        path.write_bytes(pptx_bytes("two"))
        assert saved.get(timeout=5) == path
        assert monitor.is_running
        assert any("SAVE_CALLBACK_FAILED" in r.message for r in caplog.records)
    finally:
        monitor.stop()
