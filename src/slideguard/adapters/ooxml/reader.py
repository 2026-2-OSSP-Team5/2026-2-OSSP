from hashlib import sha256
from pathlib import Path

from pptx import Presentation

from slideguard.domain.models import Box, DeckSnapshot, Shape


class OoxmlDeckReader:
    """Read object metadata without readability judgments or format inference."""

    def read(self, path: Path) -> DeckSnapshot:
        path = path.resolve(strict=True)
        if path.suffix.lower() != ".pptx" or path.name.startswith("~$"):
            raise ValueError("Choose a saved .pptx file")
        digest = sha256(path.read_bytes()).hexdigest()
        presentation = Presentation(str(path))
        shapes = tuple(
            Shape(
                slide_number=number,
                shape_id=shape.shape_id,
                name=shape.name,
                box=Box(shape.left.pt, shape.top.pt, shape.width.pt, shape.height.pt),
                text=shape.text if shape.has_text_frame else "",
                supported=False,
                notes=("Object metadata only; readability is not evaluated",),
            )
            for number, slide in enumerate(presentation.slides, 1)
            for shape in slide.shapes
        )
        if digest != sha256(path.read_bytes()).hexdigest():
            raise ValueError("File changed while reading; save and retry")
        return DeckSnapshot(
            str(path),
            digest,
            len(presentation.slides),
            presentation.slide_width.pt,
            presentation.slide_height.pt,
            shapes,
            ("No rules registered", "Effective formatting and rendered text bounds are not measured"),
        )
